# -*- coding: utf-8 -*-
"""SOP.md -> SOP.docx. python sop-docx.py [--force]

Dipakai buat menyerahkan SOP ke pengurus dalam bentuk Word. Cuma menangani konstruksi
yang benar-benar dipakai SOP.md: judul, paragraf, tabel, daftar bernomor/butir, kutipan,
**tebal**, dan `kode`. Bukan konverter markdown umum -- kalau SOP.md nanti memakai
konstruksi baru, tambahkan di sini, jangan pasang pandoc.

MENOLAK menimpa SOP.docx yang sudah ada tanpa --force: begitu berkasnya disunting di Word,
menimpanya = kehilangan pekerjaan orang.
"""
import io, os, re, sys

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

SRC, OUT = 'SOP.md', 'SOP.docx'


def runs(par, teks):
    """Tulis teks ke paragraf, hormati **tebal** dan `kode`."""
    for bagian in re.split(r'(\*\*[^*]+\*\*|`[^`]+`)', teks):
        if not bagian:
            continue
        if bagian.startswith('**') and bagian.endswith('**'):
            par.add_run(bagian[2:-2]).bold = True
        elif bagian.startswith('`') and bagian.endswith('`'):
            r = par.add_run(bagian[1:-1])
            r.font.name = 'Consolas'
            r.font.size = Pt(10)
        else:
            par.add_run(bagian)


def sel(cell, teks):
    cell.text = ''
    runs(cell.paragraphs[0], teks)


def kolom(baris):
    return [c.strip() for c in baris.strip().strip('|').split('|')]


def main(force):
    if os.path.exists(OUT) and not force:
        sys.exit(OUT + ' sudah ada. Kalau sudah disunting di Word, menimpanya menghapus '
                       'pekerjaan itu. Sengaja mau menimpa? jalankan: python sop-docx.py --force')

    baris = io.open(SRC, encoding='utf-8').read().split('\n')
    doc = Document()
    normal = doc.styles['Normal']
    normal.font.name = 'Times New Roman'
    normal.font.size = Pt(12)

    i = 0
    while i < len(baris):
        b = baris[i]

        if not b.strip() or (b.startswith('---') and set(b.strip()) == {'-'}):
            i += 1
            continue

        # tabel: kumpulkan baris berawalan |, baris pemisah dibuang
        if b.startswith('|'):
            blok = []
            while i < len(baris) and baris[i].startswith('|'):
                blok.append(baris[i])
                i += 1
            isi = [kolom(x) for x in blok if not re.match(r'^\|[\s:|-]+\|?$', x)]
            lebar = max(len(r) for r in isi)
            t = doc.add_table(rows=len(isi), cols=lebar)
            t.style = 'Table Grid'
            for r, row in enumerate(isi):
                for c in range(lebar):
                    sel(t.rows[r].cells[c], row[c] if c < len(row) else '')
                    if r == 0:
                        for p in t.rows[r].cells[c].paragraphs:
                            for run in p.runs:
                                run.bold = True
            doc.add_paragraph()
            continue

        # judul
        m = re.match(r'^(#{1,4})\s+(.*)$', b)
        if m:
            tingkat, teks = len(m.group(1)), m.group(2)
            if tingkat == 1:
                p = doc.add_heading('', 0)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p = doc.add_heading('', min(tingkat - 1, 4))
            runs(p, teks)
            i += 1
            continue

        # kutipan
        if b.startswith('> '):
            teks = []
            while i < len(baris) and baris[i].startswith('> '):
                teks.append(baris[i][2:].strip())
                i += 1
            p = doc.add_paragraph(style='Intense Quote')
            runs(p, ' '.join(teks))
            continue

        # daftar bernomor / butir. Baris menjorok = lanjutan item yang sama.
        m = re.match(r'^(\d+\.|-)\s+(.*)$', b)
        if m:
            teks = [b.strip()]
            i += 1
            while i < len(baris) and baris[i].startswith('   ') and baris[i].strip():
                teks.append(baris[i].strip())
                i += 1
            gabung = ' '.join(teks)
            if m.group(1) == '-':
                p = doc.add_paragraph(style='List Bullet')
                runs(p, gabung[2:])
            else:
                # nomor ditulis harfiah: penomoran SOP harus persis, bukan tebakan Word
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.75)
                runs(p, gabung)
            continue

        # paragraf biasa
        teks = []
        while i < len(baris) and baris[i].strip() and not re.match(r'^(\||#|>|-\s|\d+\.\s)', baris[i]):
            teks.append(baris[i].strip())
            i += 1
        p = doc.add_paragraph()
        runs(p, ' '.join(teks))

    doc.save(OUT)
    print(OUT + ' dibuat. paragraf=%d tabel=%d' % (len(doc.paragraphs), len(doc.tables)))


if __name__ == '__main__':
    main('--force' in sys.argv)
